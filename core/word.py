"""Gerador Word com identidade visual corporativa e consistente."""

from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from core.imagens import dimensoes_proporcionais
from core.estrutura import preparar_conteudo


NAVY, BLUE, SLATE, MUTED = "0F172A", "2563EB", "334155", "64748B"
LIGHT, LINE, CALLOUT, USABLE_WIDTH = (
    "F1F5F9",
    "CBD5E1",
    "EFF6FF",
    Inches(6.77),
)


def gerar_docx(metadados, secoes, caminho_saida):
    """Gera o arquivo DOCX a partir dos blocos normalizados."""
    doc = docx.Document()

    _configurar_documento(doc, metadados)

    conteudo, sumario, referencias = preparar_conteudo(secoes)

    _capa(doc, metadados)
    _sumario(doc, sumario)

    for bloco in conteudo:
        _renderizar_bloco(doc, bloco)

    if referencias:
        _bibliografia(doc, referencias)

    doc.save(caminho_saida)


def _configurar_documento(doc, metadados):
    """Configura página, estilos, cabeçalho, rodapé e propriedades."""
    section = doc.sections[0]

    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)

    _configurar_estilos(doc)
    _cabecalho_rodape(section, metadados)

    doc.core_properties.title = metadados.get(
        "nome_projeto",
        "Documentação",
    )

    doc.core_properties.author = metadados.get(
        "autor",
        "",
    )


def _configurar_estilos(doc):
    """Configura os estilos utilizados na documentação."""
    normal = doc.styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(SLATE)

    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    for nome, tamanho, cor, antes, depois in (
        ("Heading 1", 17, NAVY, 18, 8),
        ("Heading 2", 13, BLUE, 13, 6),
        ("Heading 3", 11.5, SLATE, 10, 4),
    ):
        estilo = doc.styles[nome]

        estilo.font.name = "Aptos Display"
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string(cor)

        estilo.paragraph_format.space_before = Pt(antes)
        estilo.paragraph_format.space_after = Pt(depois)
        estilo.paragraph_format.keep_with_next = True

    _estilo(
        doc,
        "Caption",
        "Aptos",
        8.5,
        MUTED,
        3,
        11,
        WD_ALIGN_PARAGRAPH.CENTER,
        True,
    )

    _estilo(
        doc,
        "Figure Title",
        "Aptos",
        10,
        SLATE,
        8,
        4,
        WD_ALIGN_PARAGRAPH.LEFT,
    )

    _estilo(
        doc,
        "Figure Source",
        "Aptos",
        9,
        SLATE,
        3,
        11,
        WD_ALIGN_PARAGRAPH.LEFT,
    )

    _estilo(
        doc,
        "Code Block",
        "Cascadia Mono",
        8.5,
        SLATE,
        4,
        10,
        espaco=1.0,
    )

    _estilo(
        doc,
        "Callout",
        "Aptos",
        10,
        SLATE,
        0,
        9,
        espaco=1.15,
    )

    _estilo(
        doc,
        "Link",
        "Aptos",
        10,
        BLUE,
        0,
        7,
    )


def _estilo(
    doc,
    nome,
    fonte,
    tamanho,
    cor,
    antes,
    depois,
    alinhamento=None,
    italico=False,
    espaco=None,
):
    """Cria ou atualiza um estilo de parágrafo."""
    estilo = (
        doc.styles[nome]
        if nome in doc.styles
        else doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
    )

    estilo.font.name = fonte
    estilo.font.size = Pt(tamanho)
    estilo.font.color.rgb = RGBColor.from_string(cor)
    estilo.font.italic = italico

    formato = estilo.paragraph_format
    formato.space_before = Pt(antes)
    formato.space_after = Pt(depois)

    if alinhamento is not None:
        formato.alignment = alinhamento

    if espaco is not None:
        formato.line_spacing = espaco


def _cabecalho_rodape(section, metadados):
    """Configura cabeçalho e rodapé do documento."""
    header = section.header.paragraphs[0]

    header.text = metadados.get(
        "nome_projeto",
        "Documentação",
    )

    header.style = "Caption"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _borda_paragrafo(
        header,
        bottom=LINE,
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run = footer.add_run("Página ")

    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    _campo(
        footer,
        "PAGE",
    )


def _capa(doc, metadados):
    """Cria a capa da documentação."""

    # ------------------------------------------------------------------
    # Organização
    # ------------------------------------------------------------------

    organizacao = (
        metadados.get("organizacao", "").strip()
        or "INSTITUIÇÃO NÃO INFORMADA"
    )

    instituicao = doc.add_paragraph(
        organizacao.upper()
    )

    instituicao.alignment = WD_ALIGN_PARAGRAPH.CENTER
    instituicao.paragraph_format.space_after = Pt(90)

    for run in instituicao.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)

    # ------------------------------------------------------------------
    # Autor
    # ------------------------------------------------------------------

    autor_texto = (
        metadados.get("autor", "").strip()
        or "AUTOR NÃO INFORMADO"
    )

    autor = doc.add_paragraph(
        autor_texto.upper()
    )

    autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    autor.paragraph_format.space_after = Pt(92)

    for run in autor.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(SLATE)

    # ------------------------------------------------------------------
    # Título
    # ------------------------------------------------------------------

    nome_projeto = (
        metadados.get("nome_projeto", "").strip()
        or "DOCUMENTAÇÃO"
    )

    titulo = doc.add_paragraph()

    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(10)

    run = titulo.add_run(
        nome_projeto.upper()
    )

    run.font.name = "Aptos"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    # ------------------------------------------------------------------
    # Tipo do documento
    # ------------------------------------------------------------------

    tipo_documento = (
        metadados.get(
            "tipo_documento",
            "",
        ).strip()
        or "Documentação Técnica"
    )

    tipo = doc.add_paragraph(
        tipo_documento
    )

    tipo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tipo.paragraph_format.space_after = Pt(8)

    for run in tipo.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(SLATE)

    # ------------------------------------------------------------------
    # Descrição
    # ------------------------------------------------------------------

    descricao = (
        metadados.get(
            "descricao",
            "",
        ).strip()
    )

    descricao_paragrafo = doc.add_paragraph(
        descricao
    )

    descricao_paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    descricao_paragrafo.paragraph_format.space_after = Pt(80)

    for run in descricao_paragrafo.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(11)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(MUTED)

    # ------------------------------------------------------------------
    # Cidade
    # ------------------------------------------------------------------

    cidade = (
        metadados.get(
            "cidade",
            "",
        ).strip()
        or "São Paulo"
    )

    local = doc.add_paragraph(
        cidade
    )

    local.alignment = WD_ALIGN_PARAGRAPH.CENTER
    local.paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------------------
    # Ano
    # ------------------------------------------------------------------

    data_criacao = (
        metadados.get(
            "data_criacao",
            "",
        ).strip()
    )

    ano_texto = (
        data_criacao.split("/")[-1]
        if data_criacao
        else ""
    )

    ano = doc.add_paragraph(
        ano_texto
    )

    ano.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------------------------------------------
    # Estilo cidade / ano
    # ------------------------------------------------------------------

    for paragrafo in (local, ano):
        for run in paragrafo.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor.from_string(SLATE)

    doc.add_page_break()


def _sumario(doc, itens):
    """Renderiza o sumário."""
    if not itens:
        return

    doc.add_heading(
        "Sumário",
        level=1,
    )

    for item in itens:
        p = doc.add_paragraph()

        p.paragraph_format.left_indent = Cm(
            (item["nivel"] - 1) * 0.45
        )

        p.paragraph_format.space_after = Pt(3)

        run = p.add_run(
            f"{item['numero']} {item['texto']}"
        )

        run.font.name = "Aptos"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(
            NAVY if item["nivel"] == 1 else SLATE
        )

        if item["nivel"] == 1:
            run.bold = True

    doc.add_page_break()


def _bibliografia(doc, referencias):
    """Renderiza as referências."""
    doc.add_heading(
        "Referências",
        level=1,
    )

    for indice, referencia in enumerate(
        referencias,
        start=1,
    ):
        p = doc.add_paragraph(
            style="Normal"
        )

        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)

        p.add_run(
            f"{indice}. "
        ).bold = True

        p.add_run(
            referencia["titulo"]
        )

        if referencia.get("autor"):
            p.add_run(
                f". {referencia['autor']}"
            )

        if referencia.get("url"):
            run = p.add_run(
                f". Disponível em: {referencia['url']}"
            )

            run.font.color.rgb = RGBColor.from_string(
                BLUE
            )

        if referencia.get("data_acesso"):
            p.add_run(
                f". Acesso em: {referencia['data_acesso']}"
            )


def _renderizar_bloco(doc, bloco):
    """Renderiza um bloco normalizado no documento Word."""
    tipo = bloco["tipo"]

    if tipo in {
        "titulo",
        "subtitulo",
        "secao_nivel_3",
    }:
        doc.add_heading(
            bloco["texto"],
            level={
                "titulo": 1,
                "subtitulo": 2,
                "secao_nivel_3": 3,
            }[tipo],
        )

    elif tipo == "texto":
        doc.add_paragraph(
            bloco["texto"]
        )

    elif tipo == "codigo":
        p = doc.add_paragraph(
            bloco["texto"],
            style="Code Block",
        )

        _sombrear_paragrafo(
            p,
            LIGHT,
        )

        _borda_paragrafo(
            p,
            left=BLUE,
        )

    elif tipo == "lista":
        for item in bloco["itens"]:
            p = doc.add_paragraph(
                item,
                style="List Bullet",
            )

            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.35)

    elif tipo in {
        "imagem",
        "diagrama_mermaid",
        "diagrama_bpmn",
    }:
        _imagem(
            doc,
            bloco,
        )

    elif tipo == "tabela":
        _tabela(
            doc,
            bloco,
        )

    elif tipo == "link":
        p = doc.add_paragraph(
            style="Link"
        )

        p.add_run(
            bloco["texto"]
        ).bold = True

        p.add_run(
            f"\n{bloco['url']}"
        )

    elif tipo == "callout":
        tabela = doc.add_table(
            rows=1,
            cols=1,
        )

        tabela.autofit = False

        _larguras(
            tabela,
            [USABLE_WIDTH],
        )

        _celula(
            tabela.cell(0, 0),
            CALLOUT,
            bloco["texto"],
            False,
            SLATE,
        )

        tabela.cell(
            0,
            0,
        ).paragraphs[0].style = "Callout"


def _tabela(doc, bloco):
    """Renderiza uma tabela."""
    if bloco.get("titulo"):
        p = doc.add_paragraph(
            bloco["titulo"],
            style="Caption",
        )

        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    colunas = bloco["cabecalhos"]

    tabela = doc.add_table(
        rows=1,
        cols=len(colunas),
    )

    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False

    _larguras(
        tabela,
        _larguras_colunas(
            len(colunas)
        ),
    )

    tabela.rows[0]._tr.get_or_add_trPr().append(
        OxmlElement("w:tblHeader")
    )

    for celula, texto in zip(
        tabela.rows[0].cells,
        colunas,
    ):
        _celula(
            celula,
            NAVY,
            texto,
            True,
            "FFFFFF",
        )

    for indice, linha in enumerate(
        bloco["linhas"]
    ):
        for celula, texto in zip(
            tabela.add_row().cells,
            linha,
        ):
            _celula(
                celula,
                "FFFFFF" if indice % 2 == 0 else LIGHT,
                texto,
                False,
                SLATE,
            )

    _bordas_tabela(
        tabela
    )


def _imagem(doc, bloco):
    """Renderiza uma imagem ou diagrama."""
    caminho = bloco.get(
        "caminho",
        "",
    )

    if not Path(caminho).is_file():
        doc.add_paragraph(
            f"[Imagem não encontrada: {caminho}]",
            style="Callout",
        )
        return

    numero = bloco.get(
        "_numero_figura",
        "",
    )

    descricao = (
        bloco.get("legenda")
        or (
            "Diagrama BPMN"
            if bloco["tipo"] == "diagrama_bpmn"
            else "Imagem"
        )
    )

    largura, _ = dimensoes_proporcionais(
        caminho,
        5.85,
    )

    p = doc.add_paragraph()

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)

    p.add_run().add_picture(
        caminho,
        width=Inches(largura),
    )

    titulo = doc.add_paragraph(
        f"Figura {numero} - {descricao}",
        style="Figure Title",
    )

    titulo.paragraph_format.keep_with_next = True

    fonte = bloco.get(
        "fonte",
        "Elaborado pelo autor.",
    )

    doc.add_paragraph(
        f"Fonte: {fonte}",
        style="Figure Source",
    )


def _larguras_colunas(quantidade):
    """Calcula as larguras das colunas da tabela."""
    if quantidade <= 0:
        return []

    if quantidade == 2:
        return [
            Inches(2.0),
            Inches(4.77),
        ]

    return [
        USABLE_WIDTH / quantidade
        for _ in range(quantidade)
    ]


def _larguras(tabela, larguras):
    """Aplica larguras às células da tabela."""
    for linha in tabela.rows:
        for celula, largura in zip(
            linha.cells,
            larguras,
        ):
            celula.width = largura


def _celula(
    celula,
    fundo,
    texto,
    negrito,
    cor,
):
    """Configura uma célula de tabela."""
    celula.text = str(texto)

    celula.vertical_alignment = (
        WD_ALIGN_VERTICAL.CENTER
    )

    _sombrear_celula(
        celula,
        fundo,
    )

    _margens_celula(
        celula
    )

    for p in celula.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)

        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.font.bold = negrito
            run.font.color.rgb = RGBColor.from_string(
                cor
            )


def _margens_celula(celula):
    """Configura as margens internas da célula."""
    tcPr = celula._tc.get_or_add_tcPr()

    margens = (
        tcPr.first_child_found_in("w:tcMar")
        or OxmlElement("w:tcMar")
    )

    for lado in (
        "top",
        "start",
        "bottom",
        "end",
    ):
        node = margens.find(
            qn(f"w:{lado}")
        )

        if node is None:
            node = OxmlElement(
                f"w:{lado}"
            )

            margens.append(
                node
            )

        node.set(
            qn("w:w"),
            "110"
            if lado in {"start", "end"}
            else "80",
        )

        node.set(
            qn("w:type"),
            "dxa",
        )

    if margens.getparent() is None:
        tcPr.append(
            margens
        )


def _sombrear_celula(celula, cor):
    """Aplica cor de fundo à célula."""
    shd = OxmlElement(
        "w:shd"
    )

    shd.set(
        qn("w:fill"),
        cor,
    )

    celula._tc.get_or_add_tcPr().append(
        shd
    )


def _sombrear_paragrafo(paragrafo, cor):
    """Aplica cor de fundo ao parágrafo."""
    shd = OxmlElement(
        "w:shd"
    )

    shd.set(
        qn("w:fill"),
        cor,
    )

    paragrafo._p.get_or_add_pPr().append(
        shd
    )


def _borda_paragrafo(paragrafo, **lados):
    """Adiciona bordas ao parágrafo."""
    bordas = OxmlElement(
        "w:pBdr"
    )

    for lado, cor in lados.items():
        borda = OxmlElement(
            f"w:{lado}"
        )

        borda.set(
            qn("w:val"),
            "single",
        )

        borda.set(
            qn("w:sz"),
            "8",
        )

        borda.set(
            qn("w:color"),
            cor,
        )

        bordas.append(
            borda
        )

    paragrafo._p.get_or_add_pPr().append(
        bordas
    )


def _bordas_tabela(tabela):
    """Adiciona bordas à tabela."""
    bordas = OxmlElement(
        "w:tblBorders"
    )

    for lado in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        borda = OxmlElement(
            f"w:{lado}"
        )

        borda.set(
            qn("w:val"),
            "single",
        )

        borda.set(
            qn("w:sz"),
            "4",
        )

        borda.set(
            qn("w:color"),
            LINE,
        )

        bordas.append(
            borda
        )

    tabela._tbl.tblPr.append(
        bordas
    )


def _campo(paragrafo, instrucao):
    """Insere um campo Word, como o número da página."""
    run = paragrafo.add_run()

    begin = OxmlElement(
        "w:fldChar"
    )
    begin.set(
        qn("w:fldCharType"),
        "begin",
    )

    instr = OxmlElement(
        "w:instrText"
    )
    instr.set(
        qn("xml:space"),
        "preserve",
    )
    instr.text = instrucao

    end = OxmlElement(
        "w:fldChar"
    )
    end.set(
        qn("w:fldCharType"),
        "end",
    )

    run._r.extend(
        (
            begin,
            instr,
            end,
        )
    )